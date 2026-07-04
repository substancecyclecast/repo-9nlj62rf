"""REST endpoints для лотов. WebSocket-роуты — в websocket.py."""
from __future__ import annotations

import tempfile
from io import BytesIO
from typing import Any
from uuid import UUID

from fastapi import APIRouter, BackgroundTasks, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from ...agents import run_lot
from ...db.models import Customer, LotStatus
from ...db.repositories import CustomerRepo, LotRepo
from ...db.session import get_session
from ...parsers import extract_text_from_path
from ..auth import Principal, get_principal
from ..metrics import LOT_CREATED, LOT_FINISHED, LOT_RUN_LATENCY
from ..schemas.lot import LotApproveIn, LotCreateIn, LotDetailOut, LotListItem, LotStatusEnum, PaginatedLots
from ..ws import bus

router = APIRouter(prefix="/lots", tags=["lots"])


async def _run_lot_bg(lot_id: str, raw_request: str, customer_id: str, phase: str) -> None:
    import time

    state: dict[str, Any] = {
        "lot_id": lot_id,
        "customer_id": customer_id,
        "raw_request": raw_request,
        "phase": phase,
        "status": "draft",
    }
    started = time.time()
    try:
        result = await run_lot(state)  # type: ignore[arg-type]
        final_status = (
            result.get("status") if isinstance(result, dict) else state.get("status", "unknown")
        )
        LOT_FINISHED.labels(status=str(final_status)).inc()
        LOT_RUN_LATENCY.observe(time.time() - started)
        await bus.publish(lot_id, {"type": "lot.finished", "status": final_status})
    except Exception as e:
        LOT_FINISHED.labels(status="failed").inc()
        LOT_RUN_LATENCY.observe(time.time() - started)
        await bus.publish(lot_id, {"type": "lot.failed", "error": str(e)})
        raise


@router.post("", response_model=dict)
async def create_lot(
    payload: LotCreateIn,
    background: BackgroundTasks,
    session: AsyncSession = Depends(get_session),
    principal: Principal = Depends(get_principal),
):
    # Multi-tenant: обычный user привязан к своему customer
    if principal.kind == "user" and principal.customer_id is not None:
        customer_obj = await session.get(Customer, str(principal.customer_id))
        customer = customer_obj or await CustomerRepo(session).get_or_create(payload.customer_name)
    else:
        customer = await CustomerRepo(session).get_or_create(payload.customer_name)
    lot = await LotRepo(session).create(
        {
            "customer_id": customer.id,
            "raw_request": payload.raw_request,
            "requester_email": payload.requester_email,
            "phase": payload.phase,
            "total_estimated_rub": payload.total_estimated_rub,
            "external_ref": payload.external_ref,
        }
    )
    LOT_CREATED.labels(customer_id=str(customer.id)).inc()
    background.add_task(_run_lot_bg, str(lot.id), payload.raw_request, str(customer.id), payload.phase)
    return {"id": str(lot.id), "status": lot.status.value}


# Регистрируем и без слэша, и со слэшем: реверс-прокси (Vercel) не проксирует
# путь с финальным слэшем, поэтому фронт обращается к /lots без слэша.
@router.get("", response_model=PaginatedLots)
@router.get("/", response_model=PaginatedLots)
async def list_lots(
    session: AsyncSession = Depends(get_session),
    principal: Principal = Depends(get_principal),
    status_filter: str | None = Query(None, alias="status"),
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
):
    status_val: LotStatus | None = None
    if status_filter:
        try:
            status_val = LotStatus(status_filter)
        except ValueError:
            status_val = None
    cust_id = principal.filter_customer_id()
    total = await LotRepo(session).count(customer_id=cust_id, status=status_val)
    offset = (page - 1) * page_size
    lots = await LotRepo(session).list(
        limit=page_size, offset=offset, customer_id=cust_id, status=status_val
    )
    pages = (total + page_size - 1) // page_size if total > 0 else 1
    items = [
        LotListItem(
            id=lot.id,
            status=LotStatusEnum(lot.status.value),
            category=lot.category,
            raw_request_preview=(lot.raw_request or "")[:160],
            created_at=lot.created_at,
            requires_human=lot.requires_human or False,
        )
        for lot in lots
    ]
    return PaginatedLots(items=items, total=total, page=page, page_size=page_size, pages=pages)


def _check_lot_access(principal: Principal, lot) -> None:
    """Raise 403, если user смотрит чужой customer's lot."""
    cust_filter = principal.filter_customer_id()
    if cust_filter is None:
        return
    if str(lot.customer_id) != str(cust_filter):
        raise HTTPException(403, "Forbidden: lot belongs to another tenant")


@router.get("/{lot_id}", response_model=LotDetailOut)
async def get_lot(
    lot_id: UUID,
    session: AsyncSession = Depends(get_session),
    principal: Principal = Depends(get_principal),
):
    lot = await LotRepo(session).get(lot_id)
    if not lot:
        raise HTTPException(404, "Lot not found")
    _check_lot_access(principal, lot)
    return LotDetailOut(
        id=lot.id,
        status=LotStatusEnum(lot.status.value),
        phase=lot.phase.value if lot.phase else None,
        category=lot.category,
        raw_request=lot.raw_request,
        parsed_items=lot.parsed_items or [],
        final_report=lot.final_report or None,
        requires_human=lot.requires_human or False,
        escalation_reason=lot.escalation_reason,
        created_at=lot.created_at,
        updated_at=lot.updated_at,
    )


@router.post("/{lot_id}/approve", response_model=dict)
async def approve_lot(
    lot_id: UUID,
    body: LotApproveIn,
    session: AsyncSession = Depends(get_session),
    principal: Principal = Depends(get_principal),
):
    repo = LotRepo(session)
    lot = await repo.get(lot_id)
    if not lot:
        raise HTTPException(404, "Lot not found")
    _check_lot_access(principal, lot)
    if principal.kind == "user" and principal.role == "viewer":
        raise HTTPException(403, "Viewer cannot approve")
    # Лимит одобрения по роли — buyer может одобрять до approval_limit_rub
    if principal.kind == "user" and principal.role == "buyer":
        from ...db.repositories import UserRepo
        if principal.user_id is not None:
            buyer = await UserRepo(session).get(principal.user_id)
            if buyer and buyer.approval_limit_rub is not None and lot.total_estimated_rub is not None:
                if float(lot.total_estimated_rub) > float(buyer.approval_limit_rub):
                    raise HTTPException(
                        403, f"Lot total {lot.total_estimated_rub} exceeds your limit"
                    )
    await repo.approve(lot_id, body.supplier_id, body.reviewer_name)
    await bus.publish(str(lot_id), {"type": "lot.approved", "supplier_id": str(body.supplier_id) if body.supplier_id else None})
    return {"status": "approved"}





@router.post("/{lot_id}/reject", response_model=dict)
async def reject_lot(
    lot_id: UUID,
    body: LotApproveIn,
    session: AsyncSession = Depends(get_session),
    principal: Principal = Depends(get_principal),
):
    repo = LotRepo(session)
    lot = await repo.get(lot_id)
    if not lot:
        raise HTTPException(404, "Lot not found")
    _check_lot_access(principal, lot)
    if principal.kind == "user" and principal.role == "viewer":
        raise HTTPException(403, "Viewer cannot reject")
    await repo.update_status(
        lot_id,
        LotStatus.rejected,
        escalation_reason=body.reviewer_name or "rejected_by_human",
    )
    await bus.publish(str(lot_id), {"type": "lot.rejected"})
    return {"status": "rejected"}


@router.post("/import-excel", response_model=dict)
async def import_excel(
    background: BackgroundTasks,
    file: UploadFile = File(...),
    customer_name: str = "SIBUR Demo",
    phase: str = "pre_nmck",
    session: AsyncSession = Depends(get_session),
):
    """Загружает Excel со списком позиций и создаёт лот.

    Ожидаемые колонки (первая строка — заголовки):
        наименование | количество | ед.изм | срок_дней | примечание
    """
    content = await file.read()
    if len(content) > 5 * 1024 * 1024:  # 5MB лимит
        raise HTTPException(413, "File too large (max 5MB)")
    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    except Exception as e:
        raise HTTPException(400, f"Cannot parse Excel: {e}")

    if len(rows) < 2:
        raise HTTPException(400, "Empty Excel (no data rows)")

    lines: list[str] = []
    for row in rows[1:]:
        if not row or not row[0]:
            continue
        name = str(row[0]).strip()
        qty = row[1] if len(row) > 1 else None
        unit = row[2] if len(row) > 2 else None
        deadline = row[3] if len(row) > 3 else None
        notes = row[4] if len(row) > 4 else None
        line = f"- {name}"
        if qty:
            line += f", {qty}"
            if unit:
                line += f" {unit}"
        if deadline:
            line += f", срок {deadline} дней"
        if notes:
            line += f" ({notes})"
        lines.append(line)

    raw_request = "Заявка из Excel-файла:\n" + "\n".join(lines)

    customer = await CustomerRepo(session).get_or_create(customer_name)
    lot = await LotRepo(session).create(
        {
            "customer_id": customer.id,
            "raw_request": raw_request,
            "phase": phase,
        }
    )
    background.add_task(_run_lot_bg, str(lot.id), raw_request, str(customer.id), phase)
    return {"id": str(lot.id), "status": lot.status.value, "n_rows": len(lines)}


@router.get("/{lot_id}/report.xlsx")
async def export_report_xlsx(
    lot_id: UUID,
    session: AsyncSession = Depends(get_session),
    principal: Principal = Depends(get_principal),
):
    lot = await LotRepo(session).get(lot_id)
    if not lot:
        raise HTTPException(404, "Lot not found")
    _check_lot_access(principal, lot)
    report = lot.final_report or {}
    top_3 = report.get("top_3") or []

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.append(["SnabAgent Report"])
    ws.append([f"Лот: {lot_id}"])
    ws.append([f"Категория: {lot.category or '—'}"])
    ws.append([f"Статус: {lot.status.value}"])
    ws.append([])
    ws.append([
        "Поставщик",
        "ИНН",
        "Общая цена ₽",
        "Уверенность",
        "Срок (дн)",
        "Причины",
        "Рекомендация",
    ])
    for r in top_3:
        ws.append([
            r.get("supplier_name") or "?",
            r.get("inn") or r.get("supplier_inn") or "?",
            r.get("total_price_rub") or r.get("total_price") or 0,
            r.get("score") or r.get("confidence") or 0,
            r.get("lead_time_days") or "—",
            r.get("rationale") or "",
            r.get("recommendation") or "",
        ])
    ws.append([])
    ws.append([f"Экономия vs средней: {report.get('savings_vs_avg_pct')}%"])
    ws.append([f"Экономия в ₽: {report.get('savings_vs_avg_rub')}"])
    ws.append([])
    ws.append(["Резюме:"])
    ws.append([report.get("summary") or "—"])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"snabagent-report-{str(lot_id)[:8]}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/parse-attachment", response_model=dict)
async def parse_attachment(file: UploadFile = File(...)) -> dict:
    """Парсит PDF/DOCX/TXT и возвращает извлечённый текст."""
    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(413, "File too large")
    name = file.filename or "upload"
    suffix = "." + name.rsplit(".", 1)[-1] if "." in name else ".txt"
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as t:
        t.write(content)
        path = t.name
    try:
        text = extract_text_from_path(path)
    except Exception as e:
        raise HTTPException(400, f"Parse failed: {e}")
    return {"text": (text or "")[:50000]}
